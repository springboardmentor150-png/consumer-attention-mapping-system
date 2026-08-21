"""
Report Generator Service
Generates PDF, Excel, and CSV reports from PostgreSQL analytics data.
"""

import os
import uuid
import logging
import csv
from datetime import datetime, timezone
from typing import Optional, Dict, Any, List
from sqlalchemy.orm import Session
from sqlalchemy import func

logger = logging.getLogger(__name__)

REPORTS_DIR = os.getenv("REPORTS_DIR", "backend/uploads/reports")

# Disclaimer and branding constants
SYSTEM_NAME = "Consumer Attention Mapping System"
DISCLAIMER_TEXT = (
    "Attention metrics are ESTIMATED based on head-pose analysis "
    "(not true eye-tracking). Purchase conversion data is only included "
    "when actual POS data is provided."
)
FOOTER_TEXT = f"{SYSTEM_NAME} | Confidential | Do not share without authorization"


class ReportGenerator:
    """Generates reports in PDF, Excel, and CSV formats from real DB data."""

    def generate(
        self,
        db: Session,
        report_type: str,
        report_format: str,
        store_id: Optional[str] = None,
        parameters: Optional[Dict[str, Any]] = None
    ) -> str:
        """
        Generate a report and return the saved file path.
        """
        os.makedirs(REPORTS_DIR, exist_ok=True)
        report_id = str(uuid.uuid4())
        normalized_format = report_format.lower().strip()
        filename = f"{report_type}_{report_id}.{normalized_format}"
        file_path = os.path.join(REPORTS_DIR, filename)

        data = self._fetch_data(db, report_type, store_id, parameters)

        if normalized_format == "csv":
            return self._generate_csv(data, file_path)
        elif normalized_format in ("excel", "xlsx"):
            return self._generate_excel(data, report_type, file_path)
        elif normalized_format == "pdf":
            return self._generate_pdf(data, report_type, store_id, file_path)
        else:
            raise ValueError(f"Unsupported report format: {report_format}")

    def _fetch_data(
        self,
        db: Session,
        report_type: str,
        store_id: Optional[str],
        parameters: Optional[Dict[str, Any]]
    ) -> List[Dict[str, Any]]:
        """Fetch analytics data from PostgreSQL based on report type."""
        from ..models import (
            ShopperSession, AttentionEvent, ProductInteraction,
            ConsumerBehavior, ProductScore, Product, Shelf,
            Recommendation, Video
        )

        if report_type == "consumer_attention":
            rows: List[Dict[str, Any]] = []
            sessions_query = db.query(ShopperSession)
            if store_id:
                sessions_query = sessions_query.join(
                    Video, ShopperSession.video_id == Video.video_id
                ).filter(Video.store_id == store_id)
            sessions = sessions_query.all()

            for s in sessions:
                att_count = db.query(AttentionEvent).filter(
                    AttentionEvent.session_id == str(s.session_id)
                ).count()
                avg_dur = db.query(func.avg(AttentionEvent.duration)).filter(
                    AttentionEvent.session_id == str(s.session_id)
                ).scalar() or 0.0

                rows.append({
                    "Session ID": f"{str(s.session_id)[:8]}...",
                    "Shopper ID (Unique)": s.shopper_code or f"SHP-{s.tracker_id:03d}",
                    "Tracker ID": f"ByteTrack #{s.tracker_id}",
                    "Behavior Segment": s.behavior_segment or "Focused Buyer",
                    "Entry Time (s)": round(float(s.entry_time or 0), 2),
                    "Exit Time (s)": round(float(s.exit_time or 0), 2),
                    "Total Dwell Time (s)": round(float(s.total_dwell_time or 0), 2),
                    "Zones Visited": len(s.zones_visited or []),
                    "Attention Events": att_count,
                    "Avg Attention Duration (s)": round(float(avg_dur), 2),
                    "Note": "Attention estimated via head-pose, not true eye-tracking"
                })
            return rows

        elif report_type == "product_engagement":
            rows = []
            query = db.query(Product)
            if store_id:
                query = query.join(Shelf, Product.shelf_id == Shelf.shelf_id).filter(
                    Shelf.store_id == store_id
                )
            products = query.all()

            for product in products:
                score = db.query(ProductScore).filter(
                    ProductScore.product_id == str(product.product_id)
                ).order_by(ProductScore.calculated_at.desc()).first()

                views = db.query(AttentionEvent).filter(
                    AttentionEvent.product_id == str(product.product_id)
                ).count()

                interactions = db.query(ProductInteraction).filter(
                    ProductInteraction.product_id == str(product.product_id)
                ).count()

                rows.append({
                    "Product Name": product.product_name,
                    "Category": product.category or "N/A",
                    "Brand": product.brand or "N/A",
                    "SKU": product.sku or "N/A",
                    "Total Views": views,
                    "Total Interactions": interactions,
                    "Attractiveness Score": round(float(score.attractiveness_score or 0), 4) if score else "N/A",
                    "Is Partial Score": score.is_partial if score else "N/A",
                    "Score Note": score.calculation_notes if score else "Not yet scored"
                })
            return rows

        elif report_type == "shelf_performance":
            rows = []
            query = db.query(Shelf)
            if store_id:
                query = query.filter(Shelf.store_id == store_id)
            shelves = query.all()

            for shelf in shelves:
                avg_att = db.query(func.avg(AttentionEvent.duration)).filter(
                    AttentionEvent.shelf_id == str(shelf.shelf_id)
                ).scalar() or 0.0
                total_events = db.query(AttentionEvent).filter(
                    AttentionEvent.shelf_id == str(shelf.shelf_id)
                ).count()
                product_count = db.query(Product).filter(
                    Product.shelf_id == str(shelf.shelf_id)
                ).count()

                rows.append({
                    "Shelf Name": shelf.shelf_name,
                    "Category": shelf.category,
                    "Total Attention Events": total_events,
                    "Avg Attention Duration (s)": round(float(avg_att), 2),
                    "Product Count": product_count
                })
            return rows

        elif report_type == "consumer_behavior":
            rows = []
            behaviors = db.query(ConsumerBehavior).all()
            for b in behaviors:
                rows.append({
                    "Segment": b.segment or "Unknown",
                    "Zones Visited": b.zones_visited_count or 0,
                    "Products Viewed": b.products_viewed_count or 0,
                    "Interactions": b.interactions_count or 0,
                    "Total Dwell Time (s)": round(float(b.total_dwell_time or 0), 2),
                    "Movement Speed": round(float(b.movement_speed or 0), 4),
                    "Comparison Behavior": b.comparison_behavior,
                    "Segment Reason": b.segment_reason or "N/A"
                })
            return rows

        else:
            # Generic: return alerts or recommendations
            recs_query = db.query(Recommendation)
            if store_id:
                recs_query = recs_query.filter(Recommendation.store_id == store_id)
            recs = recs_query.all()
            rows = []
            for r in recs:
                rows.append({
                    "Type": r.recommendation_type,
                    "Recommendation": r.recommendation_text,
                    "Reason": r.reason or "N/A",
                    "Supporting Metric": r.supporting_metric or "N/A",
                    "Confidence": r.confidence or "N/A",
                    "Expected Impact": r.expected_impact or "N/A"
                })
            return rows

    def _generate_csv(self, data: List[Dict[str, Any]], file_path: str) -> str:
        """Write data to CSV file."""
        if not data:
            data = [{"Note": "No data available for this report"}]
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=list(data[0].keys()))
            writer.writeheader()
            writer.writerows(data)
        logger.info(f"CSV report written to {file_path}")
        return file_path

    def _generate_excel(self, data: List[Dict[str, Any]], report_type: str, file_path: str) -> str:
        """Write data to Excel file using openpyxl."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            sheet_title = report_type.replace("_", " ").title()[:31]
            ws.title = sheet_title

            if not data:
                ws["A1"] = "No data available"
                wb.save(file_path)
                return file_path

            headers = list(data[0].keys())

            # Header row styling
            header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
            thin_border = Border(
                left=Side(style="thin", color="CCCCCC"),
                right=Side(style="thin", color="CCCCCC"),
                top=Side(style="thin", color="CCCCCC"),
                bottom=Side(style="thin", color="CCCCCC")
            )
            alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

            # Data rows
            for row_idx, row in enumerate(data, 2):
                is_alt = (row_idx % 2 == 0)
                for col_idx, key in enumerate(headers, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
                    cell.alignment = Alignment(horizontal="center" if isinstance(row.get(key), (int, float)) else "left", vertical="center")
                    cell.border = thin_border
                    if is_alt:
                        cell.fill = alt_row_fill

            # Auto-fit columns with calculated width
            for col_idx, header in enumerate(headers, 1):
                col_letter = get_column_letter(col_idx)
                max_len = len(str(header))
                for row in data:
                    val_str = str(row.get(header, "") or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

            # Add report metadata sheet
            meta_ws = wb.create_sheet("Report Info")
            meta_info = [
                ("Report Type", report_type),
                ("Generated At", datetime.now(timezone.utc).isoformat()),
                ("System", SYSTEM_NAME),
                ("Disclaimer", DISCLAIMER_TEXT)
            ]
            for r_idx, (label, val) in enumerate(meta_info, 1):
                c1 = meta_ws.cell(row=r_idx, column=1, value=label)
                c2 = meta_ws.cell(row=r_idx, column=2, value=val)
                c1.font = Font(bold=True)

            meta_ws.column_dimensions["A"].width = 18
            meta_ws.column_dimensions["B"].width = 70

            wb.save(file_path)
            logger.info(f"Excel report written to {file_path}")
            return file_path

        except ImportError:
            logger.warning("openpyxl not available, falling back to CSV")
            csv_path = file_path.replace(".excel", ".csv").replace(".xlsx", ".csv")
            return self._generate_csv(data, csv_path)

    def _generate_pdf(
        self,
        data: List[Dict[str, Any]],
        report_type: str,
        store_id: Optional[str],
        file_path: str
    ) -> str:
        """Generate PDF report using ReportLab."""
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.lib import colors
            from reportlab.platypus import (
                SimpleDocTemplate, Table, TableStyle, Paragraph,
                Spacer, HRFlowable
            )

            doc = SimpleDocTemplate(
                file_path,
                pagesize=landscape(A4),
                rightMargin=1.5 * cm,
                leftMargin=1.5 * cm,
                topMargin=2 * cm,
                bottomMargin=2 * cm
            )

            styles = getSampleStyleSheet()
            story = []

            # Title styles
            title_style = ParagraphStyle(
                "CustomTitle",
                parent=styles["Title"],
                fontSize=16,
                spaceAfter=10,
                textColor=colors.HexColor("#1a1a2e")
            )
            report_title = report_type.replace("_", " ").title() + " Report"
            story.append(Paragraph(SYSTEM_NAME, title_style))
            story.append(Paragraph(report_title, styles["Heading2"]))
            story.append(Paragraph(
                f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | "
                f"Store ID: {store_id or 'All Stores'}",
                styles["Normal"]
            ))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cccccc")))
            story.append(Spacer(1, 0.4 * cm))

            # Disclaimer banner
            disclaimer_style = ParagraphStyle(
                "Disclaimer",
                parent=styles["Normal"],
                fontSize=8,
                textColor=colors.HexColor("#666666"),
                backColor=colors.HexColor("#f5f5f5"),
                borderPadding=4,
                spaceAfter=10
            )
            story.append(Paragraph(
                f"⚠ DISCLAIMER: {DISCLAIMER_TEXT}",
                disclaimer_style
            ))
            story.append(Spacer(1, 0.3 * cm))

            if not data:
                story.append(Paragraph("No data available for this report.", styles["Normal"]))
            else:
                headers = list(data[0].keys())
                header_style = ParagraphStyle(
                    "TableHeader",
                    parent=styles["Normal"],
                    fontName="Helvetica-Bold",
                    fontSize=8,
                    textColor=colors.white,
                    alignment=1
                )
                cell_style = ParagraphStyle(
                    "TableCell",
                    parent=styles["Normal"],
                    fontName="Helvetica",
                    fontSize=7,
                    alignment=1
                )

                table_data = [[Paragraph(str(h), header_style) for h in headers]]
                for row in data:
                    row_cells = []
                    for h in headers:
                        val_str = str(row.get(h, "") if row.get(h) is not None else "")
                        row_cells.append(Paragraph(val_str[:60], cell_style))
                    table_data.append(row_cells)

                total_width = landscape(A4)[0] - 3 * cm
                col_width = total_width / len(headers)
                table = Table(table_data, colWidths=[col_width] * len(headers), repeatRows=1)
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4ff")]),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.append(table)

            story.append(Spacer(1, 0.5 * cm))
            story.append(Paragraph(
                FOOTER_TEXT,
                ParagraphStyle(
                    "Footer",
                    parent=styles["Normal"],
                    fontSize=7,
                    textColor=colors.HexColor("#888888"),
                    alignment=1
                )
            ))

            doc.build(story)
            logger.info(f"PDF report written to {file_path}")
            return file_path

        except ImportError:
            logger.warning("ReportLab not available, generating CSV fallback")
            csv_path = file_path.replace(".pdf", ".csv")
            return self._generate_csv(data, csv_path)


_report_generator: Optional[ReportGenerator] = None


def get_report_generator() -> ReportGenerator:
    global _report_generator
    if _report_generator is None:
        _report_generator = ReportGenerator()
    return _report_generator

